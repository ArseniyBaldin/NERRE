
import json
import subprocess

from docx import Document
from pathlib import Path

def parse_docx_modified(path, return_as_blocks=False):
    doc = Document(path)
    blocks = []

    def add_paragraph(paragraph):
        text = paragraph.text.strip()
        if not text:
            return
        style = paragraph.style.name
        if style.startswith("Heading"):
            blocks.append({"type": "heading", "text": text})
        else:
            blocks.append({"type": "paragraph", "text": text})

    def add_table(table):
        rows = []
        for row in table.rows:
            row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows.append(row_cells)
        blocks.append({"type": "table", "rows": rows})

    body = doc._body._element
    for child in body:
        if child.tag.endswith("tbl"):
            for table in doc.tables:
                if table._element == child:
                    add_table(table)
                    break
        elif child.tag.endswith("p"):
            for para in doc.paragraphs:
                if para._element == child:
                    add_paragraph(para)
                    break

    if return_as_blocks:
        return blocks
    else:
        return "\n\n".join(
            block["text"] if block["type"] != "table"
            else "\n".join(" | ".join(row) for row in block["rows"])
            for block in blocks
        )

def parse_docx(path):
    doc = Document(path)
    full_text = []

    def add_paragraph(text):
        if text and text.strip():
            full_text.append(text.strip())

    def add_table(table):
        full_text.append("[TABLE]")
        for row in table.rows:
            full_text.append("[ROW]")
            for j, cell in enumerate(row.cells):
                col_tag = f"[COL{j + 1}]"
                cell_text = cell.text.strip().replace("\n", " ")
                if cell_text:
                    full_text.append(f"{col_tag} {cell_text}")
            full_text.append("[/ROW]")
        full_text.append("[/TABLE]")

    body = doc._body._element
    for child in body:
        if child.tag.endswith("tbl"):
            for table in doc.tables:
                if table._element == child:
                    add_table(table)
                    break
        elif child.tag.endswith("p"):
            for para in doc.paragraphs:
                if para._element == child:
                    add_paragraph(para.text)
                    break

    return "\n".join(full_text)
def convert_docx_to_text(path):
    result = subprocess.run(
        ["pandoc", path, "-t", "markdown"],
        capture_output=True,
        text=True,
        check=True
    )
    return result.stdout

if __name__ == "__main__":
    input_dir = Path("../../data/docx_files")
    output_dir = Path("../../data/docx_parsed")
    output_dir.mkdir(parents=True, exist_ok=True)

    for docx_path in input_dir.glob("*.docx_files"):
        try:
            text = parse_docx(docx_path)
            output_path = output_dir / (docx_path.stem + ".json")
            with open(output_path, "w", encoding="utf-8") as f:
                json.dump({"text": text}, f, ensure_ascii=False, indent=2)
            print(f"✅ {docx_path.name} -> {output_path.name}")
        except Exception as e:
            print(f"❌ Ошибка при обработке {docx_path.name}: {e}")
